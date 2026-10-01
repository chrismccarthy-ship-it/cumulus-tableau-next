#!/usr/bin/env python3
"""
Merge Gamma PPTX exports (same theme) into one deck: parts/part1.pptx … partN.pptx → Tableau-Next-for-Builders.pptx

Gamma exports each carry a full slide master with the same layouts, so slides from parts 2..N are
copied into part 1's package: shapes are deep-copied, and every relationship on the source slide
(images, hyperlinks, charts, notes) is re-created on the new slide.
"""
import copy
import glob
import os
import sys

from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

HERE = os.path.dirname(os.path.abspath(__file__))


def copy_slide(dst, src_slide):
    layout = dst.slide_layouts[0]
    for lo in dst.slide_layouts:  # match layout by name when possible
        if lo.name == src_slide.slide_layout.name:
            layout = lo
            break
    new = dst.slides.add_slide(layout)
    for shp in list(new.shapes):  # drop placeholder shapes the layout added
        shp._element.getparent().remove(shp._element)
    for shp in src_slide.shapes:
        new.shapes._spTree.append(copy.deepcopy(shp._element))
    # background
    if src_slide._element.cSld.bg is not None:
        new._element.cSld.insert(0, copy.deepcopy(src_slide._element.cSld.bg))
    # relationships: images, media, hyperlinks
    for rel in src_slide.part.rels.values():
        if rel.reltype == RT.SLIDE_LAYOUT:
            continue
        if rel.is_external:
            new_rid = new.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            new_rid = new.part.rels.get_or_add(rel.reltype, rel.target_part)
        if new_rid != rel.rId:
            for el in new._element.iter():
                for attr in ("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed",
                             "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}link",
                             "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"):
                    if el.get(attr) == rel.rId:
                        el.set(attr, new_rid)
    return new


def main():
    parts = sorted(glob.glob(os.path.join(HERE, "parts", "part*.pptx")), key=lambda p: int(os.path.basename(p)[4:-5]))
    if not parts:
        sys.exit("no parts/part*.pptx found")
    dst = Presentation(parts[0])
    for p in parts[1:]:
        src = Presentation(p)
        for s in src.slides:
            copy_slide(dst, s)
    out = os.path.join(HERE, "Tableau-Next-for-Builders.pptx")
    dst.save(out)
    print(f"{out}: {len(dst.slides)} slides from {len(parts)} parts")


if __name__ == "__main__":
    main()
